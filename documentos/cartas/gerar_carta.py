from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Margens
section = doc.sections[0]
section.top_margin    = Cm(3)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3)
section.right_margin  = Cm(2.5)

# Estilo base
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

def add_paragraph(text='', bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                  space_before=0, space_after=6, size=12, color=None, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if text:
        run = p.add_run(text)
        run.bold   = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
    return p

def add_run_to(para, text, bold=False, italic=False, size=12):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    return run

# ── CABEÇALHO ──────────────────────────────────────────────────────────────
add_paragraph('SENADO FEDERAL', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              size=14, space_after=2)
add_paragraph('REPÚBLICA FEDERATIVA DO BRASIL', bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=2)
add_paragraph('Gabinete do Senador Nelson Trad Filho – Mato Grosso do Sul',
              align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=12)

# Linha separadora via borda inferior
p_line = doc.add_paragraph()
p_line.paragraph_format.space_after = Pt(14)
pPr = p_line._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '6')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '1A3A6B')
pBdr.append(bottom)
pPr.append(pBdr)

# ── NÚMERO / DATA ──────────────────────────────────────────────────────────
add_paragraph('Of. Gab. Nº _____ / 2026',
              align=WD_ALIGN_PARAGRAPH.RIGHT, size=10, space_after=4)
add_paragraph('Brasília – DF, _____ de ______________ de 2026.',
              align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, space_after=14)

# ── DESTINATÁRIO ───────────────────────────────────────────────────────────
dest = [
    ('Ao', False),
    ('Ilmo. Sr. Augusto Pestana', False),
    ('Cônsul-Geral do Brasil em Xangai', False),
    ('Consulado-Geral do Brasil', False),
    ('Xangai – República Popular da China', False),
]
for txt, bold in dest:
    add_paragraph(txt, bold=bold, space_after=0)

# ── ASSUNTO ────────────────────────────────────────────────────────────────
add_paragraph('', space_before=10, space_after=4)
p_ass = doc.add_paragraph()
p_ass.paragraph_format.space_before = Pt(4)
p_ass.paragraph_format.space_after  = Pt(12)
add_run_to(p_ass, 'Assunto: ', bold=True)
add_run_to(p_ass, 'Recomendação para concessão de vistos – Grupo de empresários '
           'chineses com interesse em investimentos no Estado de Mato Grosso do Sul.')

# ── CORPO ──────────────────────────────────────────────────────────────────
add_paragraph('Senhor Cônsul-Geral,', space_before=6, space_after=8)

corpo = [
    ('O signatário desta, Nelson Trad Filho, Senador da República pelo Estado de Mato '
     'Grosso do Sul, Presidente da Comissão de Relações Exteriores e Defesa Nacional do '
     'Senado Federal e Presidente da Comissão Brasil–China do Senado Federal, vem, com '
     'todo o apreço, apresentar à consideração de Vossa Senhoria o seguinte:'),

    ('Recentemente, este Gabinete teve a honra de receber, no Estado de Mato Grosso do '
     'Sul, representantes de empresas chinesas dos setores de máquinas industriais, '
     'imobiliário e materiais de construção. Durante os encontros realizados, os '
     'referidos empresários demonstraram expressivo interesse em estabelecer parcerias '
     'comerciais e realizar investimentos de relevante impacto econômico no Estado de '
     'Mato Grosso do Sul.'),

    ('Há grande expectativa, por parte desta representação parlamentar e do Governo '
     'Estadual, de que tais iniciativas se concretizem, contribuindo para a geração de '
     'empregos, o desenvolvimento regional e o fortalecimento das relações bilaterais '
     'entre o Brasil e a República Popular da China.'),

    ('Diante do exposto, solicito, com todo o respeito, que Vossa Senhoria aprecie com '
     'benevolência a concessão dos respectivos vistos de entrada no território brasileiro '
     'às pessoas a seguir relacionadas:'),
]
for txt in corpo:
    p = add_paragraph(txt, space_before=0, space_after=8, indent=1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# ── TABELA ─────────────────────────────────────────────────────────────────
pessoas = [
    ('01', 'CHEN, Zhenyu (陈振宇)',    'EQ4377323', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Gerente Geral – Área de Máquinas'),
    ('02', 'ZHOU, Hongli (周红丽)',    'EN2639150', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Gerente Geral – Área Imobiliária'),
    ('03', 'ZHOU, Hongwei (周红伟)',   'EE5259291', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Técnico de Máquinas'),
    ('04', 'ZHOU, Zixiang (周子翔)',   'EN9626588', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Técnico de Máquinas'),
    ('05', 'LI, Haoyang (李昊阳)',     'EP3391644', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Técnico de Máquinas'),
    ('06', 'CHEN, Jinjun (陈锦军)',    'EN9784705', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Gerente – Área Imobiliária'),
    ('07', 'TAO, Dongya (陶东亚)',     'EP0310062', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Gerente – Área Imobiliária'),
    ('08', 'LIU, Weihua (刘卫华)',     'EK3376699', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Gerente Geral – Materiais de Construção'),
    ('09', 'ZHANG, Jun (张军)',        'EP0234084', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Técnico – Área Imobiliária'),
    ('10', 'JIA, Yiyu (贾义雨)',       'EP0268359', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Técnico – Área Imobiliária'),
    ('11', 'TANG, Miao (唐淼)',        'EQ1737549', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Técnico de Máquinas'),
    ('12', 'LU, Jilyu (陆纪律)',       'EQ8873553', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Técnico – Materiais de Construção'),
    ('13', 'LIU, Qing (刘清)',         'EK0182501', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Projetista – Área Imobiliária'),
    ('14', 'YANG, Qingya (杨青亚)',    'EF1507864', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Técnico de Máquinas'),
    ('15', 'LIU, Zhi (刘志)',          'E97677656', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Projetista – Área Imobiliária'),
    ('16', 'ZHAN, Ming (展明)',        'ER4956378', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Gerente de Vendas – Materiais de Construção'),
    ('17', 'YU, Yang (于阳)',          'E72922880', 'Shandong Jinfengda Technology Development Co., Ltd.',          'Técnico – Materiais de Construção'),
    ('18', 'LI, Hongxiang (李红祥)',   'EM1312509', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Responsável – Área de Máquinas'),
    ('19', 'LYU, Donglin (吕冬林)',    'EQ0595998', 'Zhong Tang Sheng Yuan International Trade (Jiangsu) Co., Ltd.', 'Responsável – Área de Máquinas'),
]

headers = ['Nº', 'Nome Completo', 'Passaporte', 'Empresa', 'Cargo']
col_widths = [Cm(1), Cm(4.5), Cm(2.8), Cm(6), Cm(3.7)]

table = doc.add_table(rows=1 + len(pessoas), cols=5)
table.style = 'Table Grid'

# Cabeçalho
hdr = table.rows[0]
for i, (h, w) in enumerate(zip(headers, col_widths)):
    cell = hdr.cells[i]
    cell.width = w
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    # Fundo azul escuro
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1A3A6B')
    tcPr.append(shd)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Linhas de dados
for row_idx, (num, nome, passap, empresa, cargo) in enumerate(pessoas):
    row = table.rows[row_idx + 1]
    dados = [num, nome, passap, empresa, cargo]
    fill = 'F2F5FB' if row_idx % 2 == 1 else 'FFFFFF'
    for col_idx, (dado, w) in enumerate(zip(dados, col_widths)):
        cell = row.cells[col_idx]
        cell.width = w
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 3 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(dado)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)
        # Fundo alternado
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

# ── FECHO ──────────────────────────────────────────────────────────────────
add_paragraph('', space_before=10, space_after=4)
for txt in [
    'Coloco o Gabinete à inteira disposição de Vossa Senhoria para quaisquer esclarecimentos adicionais que se fizerem necessários.',
    'Certo de contar com a atenção e compreensão de Vossa Senhoria, subscrevo-me com elevada estima e consideração.',
]:
    p = add_paragraph(txt, space_before=0, space_after=8, indent=1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# ── ASSINATURA ─────────────────────────────────────────────────────────────
add_paragraph('', space_before=20, space_after=2)
add_paragraph('_' * 45, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_paragraph('Nelson Trad Filho', bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
for linha in [
    'Senador da República – Mato Grosso do Sul',
    'Presidente da Comissão de Relações Exteriores e Defesa Nacional do Senado Federal',
    'Presidente da Comissão Brasil–China do Senado Federal',
]:
    add_paragraph(linha, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=2)

# ── SALVAR ─────────────────────────────────────────────────────────────────
out = '/home/user/kohladvogados/documentos/cartas/carta-recomendacao-senado-nelson-trad.docx'
doc.save(out)
print('Gerado:', out)
